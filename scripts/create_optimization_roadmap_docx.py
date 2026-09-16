from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path('/Users/chance/Documents/0704-keyboard/docs/Keyboard Vault 分阶段优化建议.docx')

NAVY = '17365D'
PALE_BLUE = 'EAF2F8'
PALE_GRAY = 'F5F7F9'
LIGHT_BORDER = 'D9D9D9'
TEXT_GRAY = RGBColor(75, 85, 99)
DOC_FONT = 'Hiragino Sans GB'


def set_run_fonts(rpr):
    fonts = rpr.get_or_add_rFonts()
    for slot in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn(f'w:{slot}'), DOC_FONT)


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for margin, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{margin}'))
        if node is None:
            node = OxmlElement(f'w:{margin}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = f'w:{edge}'
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), '6')
        node.set(qn('w:color'), LIGHT_BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('第 ')
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = 'PAGE'
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    paragraph.add_run(' 页')


def style_run(run, size=None, bold=None, color=None):
    run.font.name = DOC_FONT
    set_run_fonts(run._element.get_or_add_rPr())
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


def add_paragraph(doc, text='', bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        style_run(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        style_run(rest)
    else:
        run = p.add_run(text)
        style_run(run)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        style_run(run)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        style_run(run)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        style_run(run, size=10, bold=True, color=RGBColor(255, 255, 255))
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_fill(cell, PALE_GRAY)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, len(values) - 1) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            style_run(run, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles['Normal']
normal.font.name = DOC_FONT
set_run_fonts(normal._element.get_or_add_rPr())
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(20, 24, 31)
normal.paragraph_format.line_spacing = 1.28
normal.paragraph_format.space_after = Pt(6)

for name, size, before, after in [
    ('Title', 25, 0, 16),
    ('Heading 1', 17, 16, 8),
    ('Heading 2', 13, 12, 6),
    ('Heading 3', 11.5, 9, 4),
]:
    style = styles[name]
    style.font.name = DOC_FONT
    set_run_fonts(style._element.get_or_add_rPr())
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.paragraph_format.space_before = Pt(4)
footer_p.add_run('Keyboard Vault 优化路线建议')
add_page_number(footer.add_paragraph())
for p in footer.paragraphs:
    for run in p.runs:
        style_run(run, size=9, color=TEXT_GRAY)

# Cover
title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_before = Pt(72)
title_pr = title._p.get_or_add_pPr()
title_borders = OxmlElement('w:pBdr')
for edge in ('top', 'left', 'bottom', 'right', 'between', 'bar'):
    border = OxmlElement(f'w:{edge}')
    border.set(qn('w:val'), 'nil')
    title_borders.append(border)
title_pr.append(title_borders)
style_run(title.add_run('Keyboard Vault 分阶段优化建议'), size=25, bold=True, color=RGBColor(0, 0, 0))

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(24)
style_run(subtitle.add_run('本地优先收藏管理项目实施路线'), size=14, color=TEXT_GRAY)

add_paragraph(
    doc,
    '本文档保存当前版本确认的优化方向，并给出实施顺序、依赖关系和验收标准。当前目标是先提高本地收藏管理的稳定性与性能，再考虑外部 API、云同步和在线导入。',
)
add_paragraph(
    doc,
    '建议先完成工程结构和测试保护，再处理图片按需加载与增量读写。这样可以降低修改文件系统和缓存逻辑时的数据风险。',
)

meta = doc.add_table(rows=3, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.autofit = False
set_table_borders(meta)
for row, (label, value) in zip(meta.rows, [
    ('项目', 'Keyboard Vault'),
    ('版本基准', '图片显示与缩略图机制修复后的当前版本'),
    ('优先原则', '数据安全 运行稳定 性能提升 可维护性'),
]):
    row.cells[0].width = Inches(1.25)
    row.cells[1].width = Inches(5.65)
    row.cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    row.cells[1].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_fill(row.cells[0], PALE_BLUE)
    for cell in row.cells:
        set_cell_margins(cell)
    p0 = row.cells[0].paragraphs[0]
    p1 = row.cells[1].paragraphs[0]
    style_run(p0.add_run(label), bold=True)
    style_run(p1.add_run(value))

doc.add_page_break()

doc.add_heading('路线概览', level=1)
add_paragraph(doc, '路线分为五个阶段。前四个阶段面向近期使用，第五阶段保留为后续功能扩展。')
add_table(
    doc,
    ['阶段', '核心目标', '主要项目', '建议时机'],
    [
        ('一', '建立安全基础', '模块拆分 通知系统 CI 文件容错 ID 检测 写入队列', '立即实施'),
        ('二', '降低图片开销', '旧图缩略图 WebP 按需读取 封面焦点 孤立图片清理', '基础完成后'),
        ('三', '减少重复读写', '单条增量保存 Markdown 增量读取 安全 Blob 缓存', '图片索引稳定后'),
        ('四', '补齐管理能力', '诊断页面 当前 vault 偏好 文件历史 ZIP 回滚', '核心流程稳定后'),
        ('五', '扩展内容能力', '多图相册上传 排序 主图设置 详情轮播', '后续考虑'),
    ],
    [0.55, 1.45, 3.45, 1.35],
)

doc.add_heading('阶段一 建立安全基础', level=1)
add_paragraph(doc, '这一阶段先提供可测试、可观察的工程基础，避免后续修改图片缓存和文件读写时引入难以定位的数据问题。')

stage1 = [
    ('模块拆分', '拆分 fs.ts ItemEditor.tsx AIPanel.tsx 等大型文件，将权限 读取 写入 图片 备份和诊断分开。', 'L'),
    ('统一通知', '统一显示保存成功 权限失效 图片处理失败 外部修改和备份结果。', 'M'),
    ('持续集成', '发布前依次执行测试 静态检查和生产构建，任一失败即停止部署。', 'S'),
    ('损坏文件隔离', '单个 Markdown 解析失败时继续加载其他收藏，并报告文件名与原因。', 'M'),
    ('重复 ID 检测', '连接目录和导入备份时扫描重复 identity.id，并提供安全修复。', 'S 到 M'),
    ('统一写入队列', '保存 删除 状态修改和批量编辑通过同一队列顺序执行。', 'M'),
]
add_table(doc, ['项目', '实施内容', '工作量'], stage1, [1.3, 4.8, 0.7])

doc.add_heading('阶段一 验收标准', level=2)
add_bullets(doc, [
    '一个损坏 Markdown 不影响其他收藏加载。',
    '并发触发多个保存操作时，文件内容不会相互覆盖。',
    '错误通知包含具体文件名、操作和可执行的处理建议。',
    '每次发布自动完成测试、静态检查和生产构建。',
])

doc.add_heading('阶段二 降低图片开销', level=1)
add_paragraph(doc, '当前程序启动时仍会读取所有主图与缩略图并转换为 Data URL。即使首页使用缩略图，大图已经进入内存，因此需要继续调整读取策略。')

stage2 = [
    ('补全旧缩略图', '扫描缺失项，在浏览器本地批量生成缩略图，并显示进度与失败原因。', 'M'),
    ('WebP 缩略图', '新缩略图使用 WebP，保留 JPEG 读取兼容。', 'S 到 M'),
    ('按需读取图片', '启动时只读取缩略图，打开详情或编辑器时再读取主图。', 'L'),
    ('封面焦点', '保存图片裁剪位置，避免 object cover 截断键盘主体。', 'M'),
    ('孤立图片清理', '扫描主数据和历史版本引用，只清理确认未被引用的图片。', 'M'),
]
add_table(doc, ['项目', '实施内容', '工作量'], stage2, [1.3, 4.8, 0.7])

doc.add_heading('按需读取的预期收益', level=2)
add_paragraph(doc, '一张 2048 × 1536 图片解码后约占 12 MB 内存，而一张 640 × 480 缩略图约占 1.2 MB。全部首页图片改用缩略图后，启动图片读取量预计减少 70% 到 90%，首页图片解码内存预计减少 80% 到 90%。真实结果取决于收藏数量、图片尺寸和浏览器缓存。')

doc.add_heading('阶段二 验收标准', level=2)
add_bullets(doc, [
    '首页初始化不读取 assets images 中的全部主图。',
    '没有缩略图的旧收藏可批量补全，失败项不会影响其他图片。',
    '详情页打开时读取主图，关闭后按缓存规则释放。',
    '清理功能不会删除仍被当前数据或历史版本引用的图片。',
])

doc.add_heading('阶段三 减少重复读写', level=1)
add_paragraph(doc, '这一阶段解决保存一条收藏时重复扫描整个 vault 的问题，并为更大的收藏库建立增量索引。')

stage3 = [
    ('单条增量保存', '保存成功后只更新当前内存条目，不再保存前后各完整读取一次 vault。', 'M 到 L'),
    ('Markdown 增量读取', '根据路径 修改时间和文件大小复用未变化文件的解析结果。', 'L'),
    ('安全 Blob 缓存', '通过稳定缓存键 引用计数和延迟释放降低 Data URL 内存开销。', 'L'),
]
add_table(doc, ['项目', '实施内容', '工作量'], stage3, [1.3, 4.8, 0.7])

doc.add_heading('安全 Blob 缓存规则', level=2)
add_numbered(doc, [
    '缓存键包含 vault 标识、文件路径和修改时间。',
    '同一个文件的并发请求复用同一个进行中的读取任务。',
    '组件使用图片时增加引用计数，卸载时减少引用计数。',
    '引用数归零后不立即释放，由容量限制和最近最少使用策略决定淘汰。',
    '重新读取 vault 时不统一释放所有 Blob URL。',
    '仅在断开目录或淘汰无引用缓存时调用 revokeObjectURL。',
    '使用 React Strict Mode 和并发读取测试验证图片不会再次变黑。',
])

doc.add_heading('阶段三 验收标准', level=2)
add_bullets(doc, [
    '修改一条收藏不会重新解析全部 Markdown 或重新读取全部图片。',
    '未变化的 Markdown 使用缓存结果，外部修改后能够正确失效。',
    '反复打开关闭详情和连续保存后，首页图片保持可见。',
    '缓存达到上限后内存能够稳定回落。',
])

doc.add_heading('阶段四 补齐管理能力', level=1)
add_paragraph(doc, '核心读写稳定后，增加诊断、偏好、历史版本和恢复保障，使用户能够发现并处理 vault 中的长期数据问题。')

stage4 = [
    ('收藏库诊断', '统计 Markdown 图片 缩略图和目录体积，检查损坏文件 重复 ID 缺失引用和孤立资源。', 'M'),
    ('当前 vault 偏好', '读取当前连接目录中的 settings user.md，并提供图形化设置。', 'M'),
    ('单文件历史', '每次覆盖前保存旧 Markdown，支持查看和恢复，并设置版本保留规则。', 'M'),
    ('ZIP 自动回滚', '恢复写入失败时自动恢复原 vault，避免留下半恢复状态。', 'M 到 L'),
]
add_table(doc, ['项目', '实施内容', '工作量'], stage4, [1.3, 4.8, 0.7])

doc.add_heading('单文件历史保存过程', level=2)
add_numbered(doc, [
    '用户点击保存后，程序读取当前正式 Markdown。',
    '旧内容写入 .history 收藏ID 时间 item.md。',
    '新内容先写入临时文件。',
    '临时文件成功关闭后替换正式文件。',
    '写入失败时保留原文件，并通过统一通知说明原因。',
    '历史版本页面允许预览差异并恢复。',
])
add_paragraph(doc, '历史 Markdown 可能引用旧图片，因此孤立图片清理必须同时扫描 .history。建议每条收藏保留最近 20 个版本或最近 90 天，手动标记的版本不自动删除。')

doc.add_heading('阶段四 验收标准', level=2)
add_bullets(doc, [
    '诊断结果能够定位到具体文件和建议操作。',
    'AI 尚未启用时，当前 vault 的用户偏好也能正常读取和编辑。',
    '误改单条收藏后能够恢复旧 Markdown。',
    'ZIP 恢复中途失败时，原 vault 可以自动恢复。',
])

doc.add_heading('阶段五 扩展内容能力', level=1)
add_paragraph(doc, '多图相册放在后续实施。现有数据结构已经支持 gallery，但编辑和详情交互尚未完整。')
add_bullets(doc, [
    '一次上传多张图片。',
    '拖动调整图片顺序。',
    '选择主图并为主图生成缩略图。',
    '删除单张图片前检查其他引用。',
    '详情页提供键盘操作和移动端滑动浏览。',
])

doc.add_heading('建议实施顺序', level=1)
add_paragraph(doc, '以下顺序兼顾数据风险和性能收益。模块拆分应保持纯重构，不与行为修改混在同一次变更中。')
add_table(
    doc,
    ['顺序', '工作包', '完成后再开始'],
    [
        ('1', '模块拆分和现有测试补强', '统一通知和 CI'),
        ('2', '损坏文件隔离 重复 ID 检测 写入队列', '图片索引改造'),
        ('3', '旧缩略图补全 WebP 封面焦点', '首页按需读取'),
        ('4', '首页按需读取和孤立图片清理', '增量保存'),
        ('5', '单条增量保存和 Markdown 增量读取', '安全 Blob 缓存'),
        ('6', '安全 Blob 缓存', '诊断和历史版本'),
        ('7', '诊断 偏好 历史版本 ZIP 回滚', '多图相册'),
    ],
    [0.65, 3.75, 2.65],
)

doc.add_heading('近期建议范围', level=1)
add_paragraph(doc, '近期不扩大 DeepSeek 对话 识图 zFrontier 在线导入 云同步和账户系统。已有入口可以保留，但新开发时间优先投入本地文件可靠性、图片性能和诊断能力。')

doc.add_heading('第一批建议交付内容', level=2)
add_bullets(doc, [
    '完成 fs.ts 的职责拆分并保持现有行为不变。',
    '建立统一通知系统和发布 CI。',
    '增加损坏 Markdown 隔离 重复 ID 检测和写入队列。',
    '建立图片索引，为批量生成旧缩略图做准备。',
])

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
